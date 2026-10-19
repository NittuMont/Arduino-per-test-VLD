"""Editable diode offset tests, without PSU or BLE hardware."""

from unittest.mock import Mock

import openpyxl
import pytest
from PyQt5 import QtCore, QtTest, QtWidgets
from PyQt5.QtCore import QSettings

from itech_interface import gui
from itech_interface.widgets.diode_offset_input import DiodeOffsetInput
from itech_interface.widgets.test_innesco_dialog import (
    TestInnescoDialog as InnescoDialog,
)


@pytest.fixture(scope="module")
def qapp():
    app = QtWidgets.QApplication.instance()
    if app is None:
        app = QtWidgets.QApplication([])
    return app


@pytest.fixture
def settings(tmp_path):
    return QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat)


@pytest.fixture
def window(monkeypatch, qapp, settings):
    # Do not start worker threads, scan Bluetooth or open a real PSU socket.
    monkeypatch.setattr(gui, "AsyncBLEWorker", Mock)
    monkeypatch.setattr(gui.MainWindow, "_start_psu_connect", lambda self: None)
    monkeypatch.setattr(gui.MainWindow, "_start_ble_scan_with_status", lambda self: None)
    # Tests must never read or overwrite the operator's saved preference.
    monkeypatch.setattr(
        gui.QtCore, "QSettings",
        lambda *args: QSettings(settings.fileName(), QSettings.Format.IniFormat),
    )
    window = gui.MainWindow()
    qapp.processEvents()
    yield window
    window.close()
    window.deleteLater()
    qapp.sendPostedEvents(None, QtCore.QEvent.Type.DeferredDelete)


def make_dialog(window, write_to_excel=None, raw_voltage=12.55):
    ctrl = Mock()
    ctrl.measure_voltage.return_value = raw_voltage
    return InnescoDialog(
        window,
        ctrl,
        write_to_excel if write_to_excel is not None else Mock(),
        Mock(),
        Mock(),
        diode_offset_input=window.diode_offset_input,
    )


def test_default_offset_is_in_top_bar(window, settings):
    assert gui.INNESCO_DIODE_OFFSET_V == 0.0
    assert window.diode_offset_input.value() == 0.0
    assert not settings.contains(gui.DIODE_OFFSET_SETTING)
    top_bar = window.centralWidget().layout().itemAt(0).layout()
    assert any(
        top_bar.itemAt(index).widget() is window.diode_offset_input
        for index in range(top_bar.count())
    )
    label = next(
        label for label in window.findChildren(QtWidgets.QLabel)
        if label.buddy() is window.diode_offset_input
    )
    assert label.text() == DiodeOffsetInput.LABEL_TEXT
    assert label.text() == "Caduta di tensione: (con SCR e iTech= 10,55 V"


@pytest.mark.parametrize(
    "text, expected",
    [("11.25", 11.25), ("11,25", 11.25), ("0", 0.0), ("-1,25", -1.25)],
)
def test_offset_can_be_typed_with_either_decimal_separator(window, text, expected):
    editor = window.diode_offset_input
    editor.selectAll()
    QtTest.QTest.keyClicks(editor, text)
    # Valid edits apply immediately, without Enter or a focus change.
    assert editor.value() == pytest.approx(expected)


@pytest.mark.parametrize("text", ["", "-", "abc", "nan", "inf"])
def test_incomplete_or_invalid_offset_keeps_last_valid_value(window, text):
    editor = window.diode_offset_input
    editor.setValue(11.25)
    editor.lineEdit().setText(text)
    assert editor.current_value() == 11.25
    assert editor.hasAcceptableInput()


def test_main_and_test_dialog_offsets_stay_synchronized(window):
    window.diode_offset_input.setValue(11.20)
    dialog = make_dialog(window)
    assert dialog.diode_offset_input.value() == 11.20

    dialog.diode_offset_input.setValue(11.75)
    assert window.diode_offset_input.value() == 11.75

    window.diode_offset_input.setValue(10.05)
    assert dialog.diode_offset_input.value() == 10.05
    assert gui.INNESCO_DIODE_OFFSET_V == 0.0
    dialog.reject()


def test_manual_offset_is_remembered_across_windows(window, settings):
    editor = window.diode_offset_input
    editor.selectAll()
    QtTest.QTest.keyClicks(editor, "10,55")
    settings.sync()
    assert float(settings.value(gui.DIODE_OFFSET_SETTING)) == 10.55

    reopened = gui.MainWindow()
    try:
        assert reopened.diode_offset_input.value() == 10.55
    finally:
        reopened.close()
        reopened.deleteLater()


def test_dialog_edits_update_saved_default_including_zero(window, settings):
    dialog = make_dialog(window)
    dialog.diode_offset_input.setValue(11.25)
    settings.sync()
    assert float(settings.value(gui.DIODE_OFFSET_SETTING)) == 11.25
    dialog.diode_offset_input.setValue(0.0)
    dialog.reject()
    settings.sync()
    assert float(settings.value(gui.DIODE_OFFSET_SETTING)) == 0.0


@pytest.mark.parametrize("stored", ["invalid", "nan", "inf", "-inf", 10000, None])
def test_invalid_saved_offset_falls_back_to_zero(window, stored):
    window._settings.setValue(gui.DIODE_OFFSET_SETTING, stored)
    assert window._load_diode_offset() == 0.0


def test_reopened_dialogs_do_not_leave_signal_connections(window):
    source = window.diode_offset_input
    baseline = source.receivers(source.valueChanged[float])
    for value in (10.55, 11.25, 0.0):
        dialog = make_dialog(window)
        assert source.receivers(source.valueChanged[float]) == baseline + 1
        dialog.diode_offset_input.setValue(value)
        dialog.reject()
        dialog.accept()  # Cleanup is idempotent if completion is signalled twice.
        assert source.receivers(source.valueChanged[float]) == baseline
        assert source.value() == value
    next_dialog = make_dialog(window)
    assert next_dialog.diode_offset_input.value() == 0.0
    next_dialog.reject()


def test_main_window_passes_offset_to_test_dialog(window, monkeypatch):
    window.diode_offset_input.setValue(11.25)
    monkeypatch.setattr(window, "_stop_all_tests", Mock())
    opened = []

    def inspect_dialog(dialog):
        opened.append(dialog.diode_offset_input.value())
        dialog.diode_offset_input.setValue(11.75)
        return 0

    monkeypatch.setattr(InnescoDialog, "exec_", inspect_dialog)
    window.on_test_innesco_tiristore()
    assert opened == [11.25]
    assert window.diode_offset_input.value() == 11.75


@pytest.mark.parametrize(
    "raw_voltage, offset, expected",
    [
        (12.55, 10.55, 2.00),
        (12.55, 11.25, 1.30),
        (2.34, 0.0, 2.34),
        (2.34, -0.50, 2.84),
        (12.555, 11.23, 1.32),  # Preserve the existing float/round behavior.
        (12.556, 11.23, 1.33),
    ],
)
def test_excel_uses_offset_edited_during_test(
    window, monkeypatch, tmp_path, raw_voltage, offset, expected
):
    filepath = tmp_path / "results.xlsx"
    workbook = openpyxl.Workbook()
    for name in ("Second", "Third", "Results"):
        workbook.create_sheet(name)
    workbook.worksheets[0].cell(row=2, column=1, value="7")
    workbook.save(filepath)
    workbook.close()

    window._excel_full_path = str(filepath)
    window.excel_path_edit.setText(str(filepath))
    window.matricola_edit.setText("7")
    # Avoid blocking on the success message; the real Excel writer still runs.
    monkeypatch.setattr(QtWidgets.QDialog, "exec_", lambda self: 0)
    warning = Mock()
    critical = Mock()
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", warning)
    monkeypatch.setattr(QtWidgets.QMessageBox, "critical", critical)
    writer = Mock(wraps=window._write_to_excel)
    dialog = make_dialog(window, writer, raw_voltage)
    monkeypatch.setattr(dialog, "_schedule_step", Mock())
    dialog._start_test()
    assert dialog._running

    # Change the offset after starting, without leaving the text field.
    dialog.diode_offset_input.selectAll()
    QtTest.QTest.keyClicks(dialog.diode_offset_input, f"{offset:.2f}".replace(".", ","))
    dialog._calculate_diode_drop(101.4)

    warning.assert_not_called()
    critical.assert_not_called()
    writer.assert_called_once()
    summary = writer.call_args.kwargs["summary"]
    assert f"Caduta diodo: {expected:.2f} V" in summary
    assert f"Offset diodo applicato: {offset:.2f} V" in summary
    assert window.diode_offset_input.value() == offset

    saved = openpyxl.load_workbook(filepath)
    try:
        sheet = saved.worksheets[3]
        assert sheet["H2"].value == pytest.approx(expected)
        assert sheet["J2"].value == "POS."
        assert sheet["K2"].value == 101
        assert sheet["M2"].value == "OK"
        assert sheet["N2"].value == "POS."
    finally:
        saved.close()


def test_measurement_error_does_not_write_excel(window):
    dialog = make_dialog(window)
    dialog.ctrl.measure_voltage.side_effect = OSError("measurement failed")
    dialog._calculate_diode_drop(101)
    dialog.write_to_excel.assert_not_called()
    dialog.safe_power_off.assert_called_once()


def test_help_menu_opens_operator_guide(window, monkeypatch):
    from itech_interface.widgets.help_dialog import HelpDialog

    opened = []

    def capture_dialog(dialog):
        opened.append(dialog)
        return 0

    monkeypatch.setattr(HelpDialog, "exec_", capture_dialog)
    help_menu = next(
        action.menu() for action in window.menuBar().actions()
        if action.text() == "Aiuto"
    )
    assert window._help_action in help_menu.actions()
    assert window._help_action.shortcut().toString() == "F1"
    window._help_action.trigger()

    assert len(opened) == 1
    guide = opened[0]
    assert guide.content.isReadOnly()
    assert not guide.content.openExternalLinks()
    text = guide.content.toPlainText()
    for topic in (
        gui.DEFAULT_HOST, "BLE", "matricola", "quarto foglio", "0,00 V",
        "10,55 V", "successivi avvii", "colonna H", "misura grezza",
        "100 V", "500 V", "AT+AL", "Tempo step tensione", "assenza di tensione",
    ):
        assert topic in text
    guide.accept()